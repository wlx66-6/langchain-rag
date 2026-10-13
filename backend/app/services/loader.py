"""文档解析：把上传的 PDF / Word / Excel / TXT / Markdown / CSV 提取成文字。"""

import csv
import logging
from pathlib import Path

from langchain_community.document_loaders import Docx2txtLoader, PyPDFLoader
from langchain_core.documents import Document

logger = logging.getLogger(__name__)

_ALLOWED_EXTS = {".pdf", ".docx", ".xlsx", ".xlsm", ".txt", ".md", ".markdown", ".csv"}


def allowed_ext(filename: str) -> str:
    return Path(filename).suffix.lower()


def check_allowed(filename: str) -> str:
    ext = allowed_ext(filename)
    if ext not in _ALLOWED_EXTS:
        raise ValueError(f"不支持的文件格式：{ext}（支持 PDF/Word/Excel/TXT/Markdown/CSV）")
    return ext


def load_file(path: Path, original_filename: str) -> list[Document]:
    """按扩展名分派到对应的解析器，统一返回 LangChain Document 列表。"""
    ext = check_allowed(original_filename)

    if ext == ".pdf":
        return PyPDFLoader(str(path)).load()

    if ext == ".docx":
        return Docx2txtLoader(str(path)).load()

    if ext in {".xlsx", ".xlsm"}:
        return _load_excel(path)

    if ext in {".txt", ".md", ".markdown"}:
        return _load_text(path, original_filename)

    if ext == ".csv":
        return _load_csv(path)

    return []


def _load_text(path: Path, filename: str) -> list[Document]:
    """纯文本 / Markdown：先按 UTF-8 读，失败再试 GBK，最后才降级为忽略坏字符。"""
    try:
        text = path.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        try:
            text = path.read_text(encoding="gbk")
        except UnicodeDecodeError:
            # errors="ignore" 会静默丢字，属于最后手段，必须留下日志便于排查
            text = path.read_text(encoding="utf-8", errors="ignore")
            logger.warning("文本文件解码失败，已忽略无法识别的字符 filename=%s", filename)
    if not text.strip():
        logger.warning("文本文件内容为空 filename=%s", filename)
    return [Document(page_content=text)]


def _load_excel(path: Path) -> list[Document]:
    """Excel：每个工作表读成一篇文字，单元格之间用 ' | ' 分隔。"""
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    docs: list[Document] = []
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(values_only=True):
            cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
            if cells:
                rows.append(" | ".join(cells))
        if rows:
            docs.append(Document(page_content=f"工作表「{ws.title}」：\n" + "\n".join(rows)))
    wb.close()
    if not docs:
        logger.warning("Excel 文件未解析出任何内容 filename=%s", path.name)
    return docs


def _load_csv(path: Path) -> list[Document]:
    """CSV：尝试多种常见编码，读成一篇文章。"""
    for enc in ("utf-8-sig", "utf-8", "gbk", "gb18030"):
        try:
            with open(path, newline="", encoding=enc) as f:
                reader = csv.reader(f)
                rows = [" | ".join(r) for r in reader if any(r)]
        except UnicodeDecodeError:
            continue
        if rows:
            return [Document(page_content="\n".join(rows))]
        logger.warning("CSV 文件内容为空 filename=%s", path.name)
        return []
    logger.warning("CSV 编码无法识别（已尝试 UTF-8 / GBK）filename=%s", path.name)
    return []
